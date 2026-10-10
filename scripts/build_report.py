"""Собирает короткий отчёт с титульником образца и скриншотами API."""

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUTPUT = DOCS / "Мини-отчёт_Тема_1.docx"


def add_hyperlink(paragraph, label: str, url: str) -> None:
    """Добавляет кликабельную ссылку чёрным шрифтом основного текста."""
    relationship = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    properties.append(fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_text(paragraph, text: str) -> None:
    """Переносит обычный текст и ссылки без служебной Markdown-разметки."""
    position = 0
    for match in re.finditer(r"\[([^\]]+)\]\((https?://[^)]+)\)", text):
        paragraph.add_run(text[position : match.start()].replace("`", ""))
        add_hyperlink(paragraph, match[1], match[2])
        position = match.end()
    paragraph.add_run(text[position:].replace("`", ""))


def add_body(document: Document, markdown: str) -> None:
    """Добавляет короткие разделы отчёта и два технических блока."""
    code_lines = None
    for line in markdown.splitlines():
        if line.startswith("```"):
            if code_lines is None:
                code_lines = []
            else:
                paragraph = document.add_paragraph("\n".join(code_lines))
                paragraph.paragraph_format.line_spacing = 1
                paragraph.paragraph_format.space_after = Pt(6)
                paragraph.runs[0].font.size = Pt(12)
                code_lines = None
        elif code_lines is not None:
            code_lines.append(line)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=1)
        elif line.strip() and not line.startswith("# "):
            paragraph = document.add_paragraph()
            add_text(paragraph, line)


def add_screenshots(document: Document) -> None:
    """Добавляет пояснение и краткие подписи к скриншотам проверки."""
    captions = (
        ("01_swagger.jpg", "Рисунок 1 — Маршруты API в Swagger"),
        ("02_auth_validation.jpg", "Рисунок 2 — Получение профиля: HTTP 200"),
        ("03_api_responses.jpg", "Рисунок 3 — Заглушка каталога: HTTP 501"),
    )
    for index, (filename, caption) in enumerate(captions):
        if index == 0:
            heading = document.add_heading("5 Проверка работоспособности", level=1)
            heading.paragraph_format.page_break_before = True
            document.add_paragraph(
                "Скриншоты показывают доступность Swagger, профиль с действующим "
                "access-токеном (HTTP 200) и заглушку каталога (HTTP 501)."
            )
        document.add_picture(str(DOCS / "screenshots" / filename), width=Cm(16))
        paragraph = document.paragraphs[-1]
        if index > 0:
            paragraph.paragraph_format.page_break_before = True
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.line_spacing = 1
        document.add_paragraph(caption, style="Caption")


def build_report() -> Path:
    """Сохраняет титульник и стили образца, дополняя их сокращённым содержимым."""
    document = Document(DOCS / "report_template.docx")
    cover_paragraphs = len(document.paragraphs)
    add_body(document, (DOCS / "mini_report.md").read_text(encoding="utf-8"))
    document.paragraphs[cover_paragraphs].paragraph_format.page_break_before = True
    add_screenshots(document)
    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_report())
